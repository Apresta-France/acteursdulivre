# -*- coding: utf-8 -*-
"""Convertit le classeur salons en JSON pour la migration."""
from __future__ import annotations

import json
import re
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(r"c:\Users\julie\Downloads\salons_du_livre_France_Europe_2026-2027.xlsx")
DST = Path(__file__).resolve().parents[1] / "database" / "seeds" / "salons.json"

MONTHS = {
    "janv": 1, "janvier": 1,
    "fev": 2, "fevr": 2, "fevrier": 2,
    "mars": 3,
    "avr": 4, "avril": 4,
    "mai": 5,
    "juin": 6,
    "juil": 7, "juillet": 7,
    "aout": 8,
    "sept": 9, "septembre": 9,
    "oct": 10, "octobre": 10,
    "nov": 11, "novembre": 11,
    "dec": 12, "decembre": 12,
}


def fold(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.lower()


def clean(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"non trouvé", "non trouve", "-", "via site"}:
        return ""
    return re.sub(r"\s+", " ", text)


def pick_date_blob(raw: str) -> str:
    if not raw:
        return ""
    parts = [p.strip() for p in raw.split("//") if p.strip()]
    return parts[-1] if parts else raw


def month_num(token: str) -> int | None:
    t = fold(token).replace(".", "").strip()
    for key, num in MONTHS.items():
        if t == key or t.startswith(key):
            return num
    return None


def parse_dates(raw: str) -> tuple[str | None, str | None, bool]:
    blob = pick_date_blob(raw)
    if not blob:
        return None, None, False
    folded = fold(blob)
    if "non trouve" in folded and not re.search(r"\d{4}", blob):
        return None, None, False

    confirmed = not any(
        x in folded
        for x in (
            "non trouve",
            "non confirme",
            "non encore",
            "probablement",
            "a reconfirmer",
            "dates precises",
        )
    )

    # 21 nov.-11 dec. 2026  /  8 oct.-1 nov. 2026
    m = re.search(
        r"(\d{1,2})\s*([a-zéû.]+)\.?\s*[-–]\s*(\d{1,2})\s*([a-zéû.]+)\.?\s+(\d{4})",
        folded,
    )
    if m:
        d1, mo1, d2, mo2, year = m.groups()
        a, b = month_num(mo1), month_num(mo2)
        if a and b:
            y = int(year)
            start = date(y, a, int(d1))
            end = date(y if a <= b else y + 1, b, int(d2))
            return start.isoformat(), end.isoformat(), confirmed

    # 12-18 oct. 2026
    m = re.search(r"(\d{1,2})\s*[-–]\s*(\d{1,2})\s+([a-zéû.]+)\.?\s+(\d{4})", folded)
    if m:
        d1, d2, mo, year = m.groups()
        a = month_num(mo)
        if a:
            y = int(year)
            return date(y, a, int(d1)).isoformat(), date(y, a, int(d2)).isoformat(), confirmed

    # 4 oct. 2026
    m = re.search(r"(\d{1,2})\s+([a-zéû.]+)\.?\s+(\d{4})", folded)
    if m:
        d1, mo, year = m.groups()
        a = month_num(mo)
        if a:
            day = date(int(year), a, int(d1))
            return day.isoformat(), day.isoformat(), confirmed

    # nov. 2026 / juin 2027
    m = re.search(r"([a-zéû.]+)\.?\s+(\d{4})", folded)
    if m:
        mo, year = m.groups()
        a = month_num(mo)
        if a:
            start = date(int(year), a, 1)
            return start.isoformat(), None, False

    m = re.search(r"(\d{4})", folded)
    if m:
        return date(int(m.group(1)), 1, 1).isoformat(), None, False

    return None, None, False


def slugify(text: str) -> str:
    text = fold(text)
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text[:80] or "salon"


def website(raw: str) -> str:
    raw = clean(raw)
    if not raw:
        return ""
    low = raw.lower()
    if low in {"facebook", "instagram"} or low.startswith("facebook "):
        return ""
    if " " in raw and "http" not in low:
        raw = raw.split()[0]
    if raw.startswith("http://") or raw.startswith("https://"):
        return raw
    if "." in raw:
        return "https://" + raw.lstrip("/")
    return ""


def main() -> None:
    wb = load_workbook(SRC, data_only=True)
    ws = wb["Evenements"]
    rows = list(ws.iter_rows(values_only=True))
    items = []
    seen: dict[str, int] = {}
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        name = clean(row[0])
        dates_raw = clean(row[4])
        start, end, confirmed = parse_dates(dates_raw)
        city = clean(row[5])
        slug = slugify(f"{name}-{city}")
        if slug in seen:
            seen[slug] += 1
            slug = f"{slug}-{seen[slug]}"
        else:
            seen[slug] = 1
        items.append(
            {
                "name": name,
                "slug": slug,
                "category": clean(row[1]),
                "type_label": clean(row[2]).split(" // ")[0].strip(),
                "direct": clean(row[3]).lower().startswith("direct"),
                "dates_raw": dates_raw,
                "starts_on": start,
                "ends_on": end,
                "dates_confirmed": confirmed,
                "city": city,
                "department": clean(row[6]),
                "region": clean(row[7]),
                "country": clean(row[8]) or "France",
                "venue": clean(row[9]).split(" // ")[0].strip(),
                "website": website(str(row[10] or "")),
                "attendance": clean(row[11]),
                "exhibitors": clean(row[12]),
                "ticket": clean(row[13]),
                "organizer": clean(row[14]).split(" // ")[0].strip(),
                "contact": clean(row[15]),
                "socials": clean(row[16]),
                "description": clean(row[17]).split(" // ")[0].strip(),
                "audience": clean(row[18]).split(" // ")[0].strip(),
                "notes": clean(row[19]),
            }
        )
    DST.parent.mkdir(parents=True, exist_ok=True)
    DST.write_text(json.dumps(items, ensure_ascii=False, indent=0), encoding="utf-8")
    dated = sum(1 for i in items if i["starts_on"])
    print(f"{len(items)} salons, {dated} avec date, -> {DST}")


if __name__ == "__main__":
    main()
